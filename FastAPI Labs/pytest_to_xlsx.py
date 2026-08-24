"""pytest_to_xlsx.py — convert a pytest JUnit XML report to .xlsx.

Usage: python pytest_to_xlsx.py <junit-xml> <output.xlsx>
"""
import sys
import xml.etree.ElementTree as ET

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
STATUS_FILL = {
    "PASSED": PatternFill("solid", fgColor="C6EFCE"),
    "FAILED": PatternFill("solid", fgColor="FFC7CE"),
    "ERROR": PatternFill("solid", fgColor="FFC7CE"),
    "SKIPPED": PatternFill("solid", fgColor="FFEB9C"),
}


def status_of(testcase):
    if testcase.find("failure") is not None:
        return "FAILED"
    if testcase.find("error") is not None:
        return "ERROR"
    if testcase.find("skipped") is not None:
        return "SKIPPED"
    return "PASSED"


def message_of(testcase):
    for tag in ("failure", "error"):
        el = testcase.find(tag)
        if el is not None:
            first = (el.get("message") or el.text or "").strip()
            return first.splitlines()[0] if first else "(no message)"
    return ""


def autosize(ws):
    for col in range(1, ws.max_column + 1):
        width = max(len(str(ws.cell(row=r, column=col).value or ""))
                    for r in range(1, ws.max_row + 1))
        ws.column_dimensions[get_column_letter(col)].width = min(width + 2, 60)


def main(junit_path, out_path):
    root = ET.parse(junit_path).getroot()
    suites = list(root.iter("testsuite"))

    wb = Workbook()
    ws = wb.active
    ws.title = "Test Results"
    headers = ["Test ID", "Test Name", "File", "Class",
               "Status", "Duration (s)", "Failure Message"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    ws.freeze_panes = "A2"

    for suite in suites:
        for tc in suite.iter("testcase"):
            classname = tc.get("classname", "")
            parts = classname.split(".")
            file_, klass = parts[0], ".".join(parts[1:])
            status = status_of(tc)
            row = ws.max_row + 1
            ws.append([f"{classname}::{tc.get('name')}", tc.get("name"),
                       file_, klass, status,
                       float(tc.get("time", 0) or 0), message_of(tc)])
            ws.cell(row=row, column=5).fill = STATUS_FILL.get(status, HEADER_FILL)
    autosize(ws)

    total = sum(int(s.get("tests", 0)) for s in suites)
    failures = sum(int(s.get("failures", 0)) for s in suites)
    errors = sum(int(s.get("errors", 0)) for s in suites)
    skipped = sum(int(s.get("skipped", 0)) for s in suites)
    passed = total - failures - errors - skipped

    ws2 = wb.create_sheet("Run Summary")
    rows = [
        ("Total tests", total),
        ("Passed", passed),
        ("Failed", failures),
        ("Errors", errors),
        ("Skipped", skipped),
        ("Pass rate (%)", round(passed / total * 100, 2) if total else 0),
        ("Total duration (s)", round(sum(float(s.get("time", 0) or 0) for s in suites), 3)),
        ("Timestamp", suites[0].get("timestamp", "") if suites else ""),
    ]
    ws2.append(["Key", "Value"])
    for cell in ws2[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    for key, value in rows:
        ws2.append([key, value])
    autosize(ws2)

    wb.save(out_path)
    print(f"Wrote {out_path}")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
